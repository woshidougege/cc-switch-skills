# -*- coding: utf-8 -*-
"""把工作内容填入固定格式的 Excel 工时表。

只写单元格的值，不改动任何格式（合并区域、列宽、行高、字体、公式全部保留）。

内容 JSON 格式（键为当月日期数字）：
    {"1": "……", "2": "……", "5": "……"}

用法：
    python fill_timesheet.py --template 模板.xlsx --output 输出.xlsx --year 2026 --month 9 \
        --name 张三 --entry 2026年4月27日 --project 联想 --customer 联想 \
        --content content.json --sign-date 2026/10/8
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl

HEADER_ROWS = 5          # 表头行，日期从第 6 行开始（行号 = 5 + 日期）
REMARK_COL = 5           # 备注列，一律不填
CHARGE_COL = 4           # 工时列
WORK_COL = 2             # 工作内容列


def load_content(path):
    """读取内容 JSON，返回 {日期(int): 文案}。"""
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    out = {}
    for key, value in data.items():
        day = int(key)
        if not 1 <= day <= 31:
            raise ValueError(f"日期超出范围: {day}")
        text = str(value).strip()
        if text:
            out[day] = text
    return out


def fill(template, output, year, month, name, entry, project, customer,
         content, hours, sign_date):
    wb = openpyxl.load_workbook(template)
    ws = wb["timesheet"]

    # ---- 抬头信息 ----
    ws.cell(row=1, column=1).value = f"Activity Report\n工   作   日   志                 {year}年{month}月"
    ws.cell(row=2, column=2).value = name
    ws.cell(row=2, column=6).value = entry
    ws.cell(row=3, column=2).value = project
    ws.cell(row=4, column=2).value = customer
    ws.cell(row=4, column=4).value = f"{year}年{month}月"

    # ---- 先清空，避免模板残留示例内容 ----
    for day in range(1, 32):
        row = HEADER_ROWS + day
        ws.cell(row=row, column=WORK_COL).value = None
        ws.cell(row=row, column=CHARGE_COL).value = None
        ws.cell(row=row, column=REMARK_COL).value = None

    # ---- 填内容 ----
    for day, text in sorted(content.items()):
        row = HEADER_ROWS + day
        ws.cell(row=row, column=WORK_COL).value = text
        ws.cell(row=row, column=CHARGE_COL).value = hours

    if sign_date:
        ws.cell(row=40, column=2).value = sign_date

    wb.save(output)
    return len(content)


def verify(template, output, filled_days):
    """确认格式未被改动。"""
    t = openpyxl.load_workbook(template)["timesheet"]
    o = openpyxl.load_workbook(output)["timesheet"]

    checks = {
        "合并单元格": t.merged_cells.ranges == o.merged_cells.ranges,
        "列宽": {k: v.width for k, v in t.column_dimensions.items()} ==
                {k: v.width for k, v in o.column_dimensions.items()},
        "行高": {k: v.height for k, v in t.row_dimensions.items()} ==
                {k: v.height for k, v in o.row_dimensions.items()},
        "日期列": all(t.cell(row=HEADER_ROWS + d, column=1).value ==
                      o.cell(row=HEADER_ROWS + d, column=1).value for d in range(1, 32)),
        "合计公式": o.cell(row=37, column=CHARGE_COL).value == "=SUM(D6:D36)",
        "天数公式": o.cell(row=38, column=CHARGE_COL).value == "=D37/8",
        "备注列留空": all(o.cell(row=HEADER_ROWS + d, column=REMARK_COL).value is None
                          for d in range(1, 32)),
    }

    print("-" * 46)
    print("格式校验")
    for label, ok in checks.items():
        print(f"  {'OK  ' if ok else 'FAIL'} {label}")

    missing = [d for d in range(1, 32)
               if o.cell(row=HEADER_ROWS + d, column=WORK_COL).value
               and not o.cell(row=HEADER_ROWS + d, column=CHARGE_COL).value]
    if missing:
        print(f"  FAIL 有内容但未填工时: {missing}")

    print(f"  已填天数: {filled_days}")
    return all(checks.values()) and not missing


def main():
    ap = argparse.ArgumentParser(description="把工作内容填入 Excel 工时表")
    ap.add_argument("--template", required=True, help="模板 xlsx 路径")
    ap.add_argument("--output", required=True, help="输出 xlsx 路径")
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--month", type=int, required=True)
    ap.add_argument("--content", required=True, help="内容 JSON 路径")
    ap.add_argument("--name", default="", help="人员姓名")
    ap.add_argument("--entry", default="", help="进场时间")
    ap.add_argument("--project", default="", help="项目名称")
    ap.add_argument("--customer", default="", help="客户")
    ap.add_argument("--hours", type=float, default=8, help="每天工时，默认 8")
    ap.add_argument("--sign-date", default="", help="签名日期，留空则不填")
    args = ap.parse_args()

    content = load_content(args.content)
    if not content:
        print("内容为空，未写入", file=sys.stderr)
        return 1

    hours = int(args.hours) if float(args.hours).is_integer() else args.hours
    filled = fill(args.template, args.output, args.year, args.month,
                  args.name, args.entry, args.project, args.customer,
                  content, hours, args.sign_date)
    print(f"已写入 {filled} 天 -> {args.output}")
    return 0 if verify(args.template, args.output, filled) else 1


if __name__ == "__main__":
    sys.exit(main())